"""Create the two-page Vietnamese functional description for the project."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[2] / 'Mo_ta_chuc_nang_website_quan_ly_lop_hoc.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.7)
section.footer_distance = Cm(0.8)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor.from_string('24343B')
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    doc.styles[style_name].font.name = 'Times New Roman'
    doc.styles[style_name]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')

heading = doc.styles['Heading 2']
heading.font.size = Pt(12.5)
heading.font.bold = True
heading.font.color.rgb = RGBColor.from_string('176D61')
heading.paragraph_format.space_before = Pt(9)
heading.paragraph_format.space_after = Pt(4)
heading.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run('CLASSROOM  |  HỆ THỐNG QUẢN LÝ LỚP HỌC THÔNG MINH')
run.font.name = 'Times New Roman'
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string('7C8E91')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('Mô tả chức năng  •  Trang ')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor.from_string('7C8E91')
for instruction in ['PAGE', 'NUMPAGES']:
    if instruction == 'NUMPAGES':
        footer.add_run(' / ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    footer._p.append(field)

def paragraph(text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def heading2(text):
    doc.add_paragraph(text, style='Heading 2')

def lead(label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(label + ' ').bold = True
    p.add_run(text)
    return p

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(5)
title.paragraph_format.keep_with_next = True
r = title.add_run('MÔ TẢ CHỨC NĂNG WEBSITE\nQUẢN LÝ LỚP HỌC THÔNG MINH')
r.bold = True
r.font.size = Pt(17)
r.font.color.rgb = RGBColor.from_string('164F48')
subtitle = doc.add_paragraph('Nền tảng web kết hợp ESP32 DevKit V1, RFID RC522 và DHT22')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(10)
subtitle.runs[0].italic = True
subtitle.runs[0].font.size = Pt(11)

heading2('1. Mục đích và phạm vi hệ thống')
paragraph('Website hỗ trợ giáo viên và người quản lý theo dõi sĩ số, quản lý học sinh, tra cứu lịch sử điểm danh và quan sát điều kiện môi trường lớp học trên một giao diện tiếng Việt. Hệ thống kết hợp thẻ RFID để ghi nhận vào/ra lớp, cảm biến DHT22 để đo nhiệt độ, độ ẩm và camera riêng để xem hình ảnh trực tiếp.')
paragraph('Phiên bản hiện tại tạo sẵn bốn lớp: 10A1, 10A2, 11A1 và 12A1. Trong đó, lớp 10A1 được triển khai đầy đủ chức năng; ba lớp còn lại là lớp mẫu, chưa kích hoạt, phục vụ định hướng mở rộng.')

heading2('2. Quản lý lớp học và hồ sơ học sinh')
paragraph('Người dùng lựa chọn lớp từ các thẻ lớp học hoặc mục Quản lý lớp học. Với lớp 10A1, website hiển thị phòng học, tổng số học sinh và các thông tin theo dõi. Khi chọn lớp chưa kích hoạt, giao diện thông báo trạng thái và cho phép quay lại lớp đang hoạt động.')
paragraph('Chức năng quản lý học sinh cho phép thêm mới, sửa họ tên, gán hoặc thay đổi mã UID thẻ RFID và tìm kiếm theo tên hoặc mã thẻ. Mỗi UID chỉ được gán cho một học sinh. Có thể lọc danh sách theo trạng thái đang trong lớp, đã rời lớp hoặc chưa đến lớp; học sinh đang trong lớp cần quét ra trước khi được xóa.')

heading2('3. Điểm danh vào/ra bằng thẻ RFID')
paragraph('Đầu đọc RC522 đọc UID trên thẻ và gửi sự kiện qua ESP32 đến máy chủ. Nếu thẻ đã được đăng ký, hệ thống xác định trạng thái hiện tại để xử lý: học sinh đang ngoài lớp quét thẻ sẽ được ghi nhận vào lớp; quét lần tiếp theo khi đang trong lớp sẽ được ghi nhận ra lớp. Sĩ số tăng hoặc giảm tương ứng và lịch sử được lưu kèm thời gian.')
lead('Chống ghi nhận trùng:', 'cùng một thẻ quét trong vòng 3 giây sẽ được bỏ qua. Khi mạng chập chờn, sự kiện gửi lại có cùng mã nhận diện không bị tính thêm lần nữa. Người dùng cần rút thẻ khỏi đầu đọc trước lượt quét tiếp theo.')
lead('Xử lý thẻ lạ:', 'thẻ chưa đăng ký được ghi vào lịch sử và hiển thị UID để quản trị viên gán cho học sinh, nhưng không làm thay đổi sĩ số.')

heading2('4. Bảng tổng quan và thống kê sĩ số')
paragraph('Trang Tổng quan hiển thị số học sinh đang trong lớp trên tổng sĩ số, tỷ lệ hiện diện, số học sinh đã đến trong ngày, số đã rời lớp và số chưa đến. Biểu đồ vòng giúp quan sát nhanh mức độ hiện diện; danh sách chi tiết thể hiện trạng thái cùng lượt quét gần nhất. Chỉ số “đã đến hôm nay” đếm học sinh có lượt vào trong ngày theo giờ Việt Nam; trạng thái đang trong lớp được giữ qua ngày và qua lần khởi động lại máy chủ.')

doc.add_page_break()
continuation = doc.add_paragraph()
continuation.paragraph_format.space_after = Pt(5)
r = continuation.add_run('CHỨC NĂNG GIÁM SÁT VÀ VẬN HÀNH')
r.bold = True
r.font.size = Pt(15)
r.font.color.rgb = RGBColor.from_string('164F48')

heading2('5. Theo dõi nhiệt độ và độ ẩm')
paragraph('DHT22 kết nối với ESP32 DevKit V1 để đo nhiệt độ phòng theo đơn vị °C và độ ẩm không khí theo đơn vị %. ESP32 gửi dữ liệu cảm biến khoảng 10 giây một lần; giao diện đồng bộ với máy chủ khoảng 2 giây một lần. Website hiển thị giá trị mới nhất và biểu đồ xu hướng, hỗ trợ theo dõi sự thay đổi của môi trường học tập.')
paragraph('Khi nhiệt độ vượt 32°C, giao diện thể hiện thông báo nhiệt độ cao. Nếu không có mẫu cảm biến mới quá 60 giây, chỉ số hiện tại chuyển sang trạng thái chưa có dữ liệu mới. Phần mềm không tự tạo giá trị thay thế khi cảm biến đọc lỗi.')

heading2('6. Quan sát lớp học bằng camera')
paragraph('Người dùng có thể chọn webcam/camera USB của máy tính hoặc nhập địa chỉ luồng hình ảnh trực tiếp từ camera IP, ESP32-CAM riêng. Với webcam, trình duyệt yêu cầu cấp quyền camera; với camera mạng, nguồn hình ảnh cần tương thích với cách hiển thị trên trình duyệt, chẳng hạn luồng MJPEG. Giao diện có chức năng mở, thay đổi nguồn và tắt camera.')
paragraph('Chức năng camera chỉ phục vụ quan sát trực tiếp, không ghi hình, không nhận diện khuôn mặt và không tự đếm học sinh. ESP32 DevKit V1 trong thiết kế này không có camera tích hợp; nhiệt độ và độ ẩm được đo bằng DHT22, không được suy ra từ hình ảnh camera.')

heading2('7. Lịch sử điểm danh, báo cáo và lưu dữ liệu')
paragraph('Mục Lịch sử điểm danh ghi nhận học sinh, UID thẻ, hoạt động và thời gian vào/ra lớp. Người dùng có thể tìm kiếm, lọc theo loại hoạt động và xem các lượt quét gần nhất; thẻ chưa đăng ký và lượt quét trùng cũng được phân biệt. Nút Xuất báo cáo tạo tệp CSV để mở bằng Excel, phục vụ kiểm tra và tổng hợp lịch sử đã lưu.')
paragraph('Ở chế độ thiết bị thật, máy chủ sử dụng SQLite để lưu hồ sơ học sinh, lịch sử và trạng thái hiện diện. Dữ liệu vẫn được giữ khi dừng rồi mở lại chương trình. Giao diện lịch sử hiển thị tối đa 500 lượt gần nhất, còn chức năng xuất báo cáo lấy toàn bộ lịch sử trên máy chủ.')

heading2('8. Kết nối thiết bị và hai chế độ sử dụng')
lead('Chế độ dùng thử:', 'cung cấp 32 học sinh mẫu cùng số liệu minh họa, cho phép trải nghiệm thao tác quản lý và quét thử thẻ. Dữ liệu dùng thử được lưu riêng trong trình duyệt, không được coi là dữ liệu cảm biến hoặc điểm danh thực tế.')
lead('Chế độ thiết bị thật:', 'ESP32 và máy tính chạy máy chủ kết nối cùng mạng Wi-Fi nội bộ. ESP32 gửi sự kiện thẻ, nhiệt độ và độ ẩm đến máy chủ; trình duyệt lấy dữ liệu để cập nhật giao diện. Danh sách học sinh thật ban đầu trống và cần được thêm, gán UID trước khi sử dụng.')
paragraph('Mục Thiết bị & kết nối hiển thị trạng thái trực tuyến/ngoại tuyến, IP của ESP32, tín hiệu Wi-Fi, lần liên lạc cuối và UID thẻ lạ gần nhất. Khóa quản trị dùng để truy cập dữ liệu; khóa thiết bị dùng để gửi dữ liệu từ ESP32. Khi mất mạng, chương trình ESP32 có hàng đợi tối đa 32 lượt quét để gửi lại; thời gian lịch sử là lúc máy chủ nhận sự kiện.')
lead('Quy trình sử dụng:', 'chạy máy chủ nội bộ → kết nối chế độ thật → thêm học sinh và gán thẻ → nạp chương trình ESP32 → quét vào/ra và theo dõi trên website. Bản trực tuyến phục vụ trải nghiệm; vận hành phần cứng dùng giao diện máy chủ nội bộ.')

p = doc.add_paragraph('Phạm vi kiểm chứng: phần mềm đã qua 18 kiểm tra tự động và chương trình ESP32 đã biên dịch thành công. Việc đọc thẻ, đo cảm biến và xem camera trên thiết bị thực tế cần được kiểm tra sau khi đấu nối.')
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(0)
for run in p.runs:
    run.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string('6F8184')

doc.core_properties.title = 'Mô tả chức năng website quản lý lớp học thông minh'
doc.core_properties.subject = 'Quản lý lớp học, điểm danh RFID và giám sát môi trường bằng ESP32'
doc.core_properties.author = 'Classroom'
doc.core_properties.keywords = 'ESP32, RC522, RFID, DHT22, quản lý lớp học'
doc.core_properties.language = 'vi-VN'
settings = doc.settings.element
update = OxmlElement('w:updateFields')
update.set(qn('w:val'), 'true')
settings.append(update)
doc.save(OUTPUT)
print(OUTPUT)
print('Word count:', sum(len(p.text.split()) for p in doc.paragraphs))
